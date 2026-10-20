from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\USUARIO\Documents\Codex\SafeWords\documentos")
OUTPUT = ROOT / "SafeWords_Documento_Wireframes_Mockups.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color="D9D9D9", size="8"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=130, bottom=110, end=130):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name="Arial", size=11, bold=False, color="000000"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, space_after=6, line_spacing=1.08):
    paragraph.paragraph_format.space_after = Pt(space_after)
    paragraph.paragraph_format.line_spacing = line_spacing


def add_body(doc, text, space_after=8):
    paragraph = doc.add_paragraph()
    style_paragraph(paragraph, space_after)
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color="243447")
    return paragraph


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.space_before = Pt(16 if level == 1 else 10)
    paragraph.paragraph_format.space_after = Pt(6)
    run = paragraph.add_run(text)
    set_run_font(run, size=16 if level == 1 else 12.5, bold=True, color="000000")
    return paragraph


def add_caption(doc, text):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(10)
    run = paragraph.add_run(text)
    set_run_font(run, size=9.5, color="64748B")
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    style_paragraph(paragraph, 3)
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color="243447")
    return paragraph


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.autofit = False
    table.allow_autofit = False
    header = table.rows[0].cells
    for index, value in enumerate(headers):
        header[index].width = Inches(widths[index])
        header[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(header[index], "173F5F")
        set_cell_borders(header[index])
        set_cell_margins(header[index])
        p = header[index].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(value)
        set_run_font(r, size=10, bold=True, color="FFFFFF")
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].width = Inches(widths[index])
            cells[index].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_borders(cells[index])
            set_cell_margins(cells[index])
            if row_index % 2 == 1:
                set_cell_shading(cells[index], "F4F8FB")
            p = cells[index].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(value)
            set_run_font(r, size=9.5, color="243447")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(11)
    for name in ("Heading 1", "Heading 2", "Title"):
        style = doc.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.color.rgb = RGBColor(0, 0, 0)


doc = Document()
configure_styles(doc)
section = doc.sections[0]
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.7)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

# Cover and introduction
title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(8)
run = title.add_run("Documento de Wireframes y Mockups")
set_run_font(run, size=25, bold=True, color="000000")

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(22)
run = subtitle.add_run("Proyecto SafeWords")
set_run_font(run, size=14, color="526579")

add_body(
    doc,
    "Este documento presenta la evolucion visual de SafeWords, una plataforma web tipo foro sobre convivencia digital, bullying, ciberbullying y violencia. Primero se muestra un wireframe de baja fidelidad, elaborado como boceto para validar la estructura. Luego se presenta un mockup de mayor fidelidad, cercano a la interfaz desarrollada con Angular y Bootstrap.",
    space_after=12,
)
add_body(
    doc,
    "La separacion entre ambas etapas permite comprobar primero la organizacion de la informacion y despues la apariencia, los colores, las acciones y la adaptacion a dispositivos moviles.",
    space_after=12,
)

add_heading(doc, "1 Objetivo del documento", 1)
add_body(doc, "Definir y comunicar la propuesta visual de las pantallas principales de SafeWords antes y despues de aplicar el diseño de interfaz.")
add_table(
    doc,
    ["Etapa", "Proposito", "Resultado"],
    [
        ["Wireframe", "Representar la estructura sin distraerse con colores o detalles visuales.", "Boceto de pantallas, jerarquia y flujo."],
        ["Mockup", "Mostrar como se veria la interfaz con la identidad visual del proyecto.", "Propuesta de colores, componentes y estados."],
    ],
    [1.0, 3.1, 2.55],
)

add_heading(doc, "2 Wireframes de baja fidelidad", 1)
add_body(doc, "El wireframe se presenta como un dibujo intencionalmente simple. Sus lineas, cajas y textos representan la ubicacion de los componentes, no el diseño final. Esta etapa permite detectar si una persona entiende el recorrido de la aplicacion antes de invertir tiempo en estilos.")
wireframe = ROOT / "SafeWords_wireframes_boceto.png"
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(wireframe), width=Inches(6.55))
add_caption(doc, "Figura 1. Wireframes iniciales para la comunidad, la publicacion y la moderacion.")

add_heading(doc, "2.1 Pantallas representadas", 2)
add_bullet(doc, "Inicio y comunidad: permite leer publicaciones, buscar por tema y consultar comentarios sin iniciar sesion.")
add_bullet(doc, "Crear publicacion: requiere una cuenta y permite elegir si el comentario o la publicacion se muestra de manera anonima.")
add_bullet(doc, "Moderacion: muestra el espacio privado del administrador para revisar reportes, aprobar, ocultar o eliminar contenido.")

doc.add_page_break()
add_heading(doc, "3 Mockups de la interfaz", 1)
add_body(doc, "El mockup utiliza la identidad visual de SafeWords: tonos azul y verde para transmitir confianza y acompañamiento, fondos claros para facilitar la lectura, tarjetas con bordes suaves y botones visibles. La propuesta mantiene la logica del frontend actual, pero la presenta de manera visual para el documento del proyecto.")
mockup = ROOT / "SafeWords_mockups_interfaz.png"
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(str(mockup), width=Inches(6.55))
add_caption(doc, "Figura 2. Mockups de alta fidelidad para la comunidad, la creacion de publicaciones y la administracion.")

add_heading(doc, "3.1 Descripcion de las pantallas", 2)
add_table(
    doc,
    ["Pantalla", "Usuarios", "Elementos principales", "Acceso"],
    [
        ["Comunidad", "Visitantes y usuarios", "Navegacion, hero, busqueda, categorias y publicaciones.", "Publica."],
        ["Crear publicacion", "Usuarios autenticados", "Titulo, contenido, selector de anonimato y boton publicar.", "Privado."],
        ["Panel de administracion", "Administrador", "Resumen, reportes pendientes, usuarios y acciones de moderacion.", "Privado y protegido."],
    ],
    [1.25, 1.25, 2.7, 1.45],
)

add_heading(doc, "4 Criterios de experiencia de usuario", 1)
add_bullet(doc, "Lectura publica: cualquier visitante puede consultar la informacion y las conversaciones sin registrarse.")
add_bullet(doc, "Participacion segura: para publicar, comentar o reportar se solicita iniciar sesion.")
add_bullet(doc, "Anonimato controlado: la comunidad ve el nombre anonimo cuando la persona lo elige; el administrador conserva los datos necesarios para moderar.")
add_bullet(doc, "Acciones claras: publicar, comentar, buscar y reportar se presentan con botones visibles y mensajes de confirmacion.")
add_bullet(doc, "Adaptacion responsive: las tarjetas se apilan, el menu se compacta y los formularios ocupan todo el ancho disponible en celulares.")

add_heading(doc, "5 Flujo visual principal", 1)
add_body(doc, "El recorrido propuesto comienza en la comunidad publica. La persona puede leer y buscar informacion. Cuando decide participar, entra a registro o inicio de sesion. Despues puede crear una publicacion, comentar o reportar contenido. Los reportes llegan al panel privado del administrador, quien puede revisar y gestionar cada caso.")
add_table(
    doc,
    ["Paso", "Accion", "Resultado"],
    [
        ["1", "Visitar SafeWords", "Se muestran publicaciones e informacion publica."],
        ["2", "Iniciar sesion", "Se habilitan publicar, comentar y reportar."],
        ["3", "Elegir anonimato", "La participacion se muestra con nombre o como anonima."],
        ["4", "Reportar contenido", "El sistema registra el reporte para moderacion."],
        ["5", "Revisar como administrador", "El contenido se aprueba, oculta o elimina segun las normas."],
    ],
    [0.55, 2.25, 3.85],
)

add_heading(doc, "6 Conclusion", 1)
add_body(doc, "Los wireframes permiten validar la estructura y el funcionamiento esperado sin depender del diseño visual. Los mockups muestran como esa estructura se transforma en una interfaz coherente con SafeWords, con una experiencia publica de lectura, participacion autenticada, anonimato configurable y moderacion segura. Ambas representaciones hacen visible la relacion entre la idea inicial y la pagina implementada.")

doc.save(OUTPUT)
print(OUTPUT)
